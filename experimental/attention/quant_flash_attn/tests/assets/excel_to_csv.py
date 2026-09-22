#!/usr/bin/python3
# -*- coding: utf-8 -*-
# -----------------------------------------------------------------------------------------------------------
# Copyright (c) 2026 Huawei Technologies Co., Ltd.
# This program is free software, you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for the full text of the License.
# -----------------------------------------------------------------------------------------------------------
"""把 mxFp4 Excel sheet 转成符合 ttk e2e 标准的 csv。

主算子直接由 ttk 调用 torch.ops.cann_ops_transformer.quant_flash_attn,
metadata 由 spec 的 npu_preprocess 在主算子调用前生成并回填 metadata slot。
输出单份 csv: qfa_mxfp4.csv (api_name=torch.ops.cann_ops_transformer.quant_flash_attn)。

tensor 顺序 (共 15 个, 对齐算子 schema 位置参数顺序):
  0  q              1  k              2  v
  3  q_descale      4  k_descale      5  v_descale
  6  block_table    7  p_scale
  8  cu_seqlens_q   9  cu_seqlens_kv  10 seqused_q  11 seqused_kv
  12 sinks          13 attn_mask      14 metadata

cu_seqlens/seqused 的真实值以 `*_values` 属性保留在 attributes, 不与算子 schema
的同名 Tensor 参数撞名 (inputs/npu_preprocess 用 `*_values` 覆盖随机占位 tensor)。
"""

import argparse
import csv
import math
import os
import sys

import openpyxl

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

API_NAME = "torch.ops.cann_ops_transformer.quant_flash_attn"

CSV_PROFILES = [
    ("qfa_mxfp4.csv", API_NAME, ""),
]

# 字段 key -> Excel 表头名 (列位置无关, 启动时按表头名动态查列号)
HEADER = {
    "name": "testcase_name",
    "custom_info": "custom_info",
    "attn_out_dtype": "attn_out_dtype",
    "q_shape": "q_shape",
    "q_dtype": "q_dtype",
    "q_datarange": "q_datarange",
    "k_shape": "k_shape",
    "k_dtype": "k_dtype",
    "k_datarange": "k_datarange",
    "v_shape": "v_shape",
    "v_dtype": "v_dtype",
    "v_datarange": "v_datarange",
    "q_descale_shape": "q_descale_shape",
    "q_descale_dtype": "q_descale_dtype",
    "q_descale_datarange": "q_descale_datarange",
    "k_descale_shape": "k_descale_shape",
    "k_descale_dtype": "k_descale_dtype",
    "k_descale_datarange": "k_descale_datarange",
    "v_descale_shape": "v_descale_shape",
    "v_descale_dtype": "v_descale_dtype",
    "v_descale_datarange": "v_descale_datarange",
    "block_table_shape": "block_table_shape",
    "block_table_dtype": "block_table_dtype",
    "p_scale_value": "p_scale_value",
    "p_scale_shape": "p_scale_shape",
    "p_scale_dtype": "p_scale_dtype",
    "p_scale_datarange": "p_scale_datarange",
    "learnable_sink_shape": "learnable_sink_shape",
    "learnable_sink_dtype": "learnable_sink_dtype",
    "attn_mask_shape": "attn_mask_shape",
    "attn_mask_dtype": "attn_mask_dtype",
    "cu_seqlens_q_value": "cu_seqlens_q_value",
    "cu_seqlens_q_shape": "cu_seqlens_q_shape",
    "cu_seqlens_q_dtype": "cu_seqlens_q_dtype",
    "cu_seqlens_kv_value": "cu_seqlens_kv_value",
    "cu_seqlens_kv_shape": "cu_seqlens_kv_shape",
    "cu_seqlens_kv_dtype": "cu_seqlens_kv_dtype",
    "seqused_q_value": "seqused_q_value",
    "seqused_q_shape": "seqused_q_shape",
    "seqused_q_dtype": "seqused_q_dtype",
    "seqused_kv_value": "seqused_kv_value",
    "seqused_kv_shape": "seqused_kv_shape",
    "seqused_kv_dtype": "seqused_kv_dtype",
    "softmax_lse_dtype": "softmax_lse_dtype",
    "quant_mode": "quant_mode",
    "softmax_scale": "softmax_scale",
    "mask_mode": "mask_mode",
    "win_left": "win_left",
    "win_right": "win_right",
    "max_seqlen_q": "max_seqlen_q",
    "max_seqlen_kv": "max_seqlen_kv",
    "layout_q": "layout_q",
    "layout_q_descale": "layout_q_descale",
    "layout_kv": "layout_kv",
    "layout_out": "layout_out",
    "return_softmax_lse": "return_softmax_lse",
    "batch_size": "batch_size",
    "num_heads_q": "num_heads_q",
    "num_heads_kv": "num_heads_kv",
    "head_dim": "head_dim",
}

OPTIONAL_HEADER = ("precision_tolerances", "absolute_precision")

DEFAULT_PRECISION_TOLERANCES = "((0.0078125, 0.005), (0.0078125, 0.005))"
DEFAULT_ABSOLUTE_PRECISION = "(0.0001, 0.0001)"

METADATA_STRIDE = 16


# -----------------------------------------------------------------------------------------------------------
# metadata slot shape 推导 (与 torch_extension/quant_flash_attn.py 一致)
# -----------------------------------------------------------------------------------------------------------
_CORE_NUMS = None


def _get_core_nums():
    global _CORE_NUMS
    if _CORE_NUMS is None:
        try:
            import torch
        except ImportError:
            torch = None
        npu = getattr(torch, "npu", None) if torch is not None else None
        if npu is None or not npu.is_available():
            _CORE_NUMS = (36, 72)
        else:
            props = npu.get_device_properties()
            _CORE_NUMS = (props.cube_core_num, props.vector_core_num)
    return _CORE_NUMS


def _calculate_max_schedule_size(batch_size, num_heads_kv, aic_num, aiv_num):
    align_size = 4096
    batch_size = batch_size if batch_size and batch_size > 0 else 1
    fa_size = aic_num * METADATA_STRIDE * batch_size * num_heads_kv
    fd_size = aiv_num * METADATA_STRIDE * batch_size * num_heads_kv
    schedule_size = METADATA_STRIDE + fa_size + fd_size
    return ((schedule_size + align_size - 1) // align_size) * align_size


def _metadata_batch_size(B, sq, skv, cu_sq, cu_skv, layout_q):
    if sq:
        return len(sq)
    if cu_sq:
        return max(len(cu_sq) - 1, 0)
    return int(B) if B else 0


def _metadata_slot_shape(B, N_kv, sq, cu_sq, layout_q):
    num_heads_kv = int(N_kv) if N_kv else 1
    batch_size = _metadata_batch_size(B, sq, [], cu_sq, [], layout_q)
    aic_num, aiv_num = _get_core_nums()
    return (2, _calculate_max_schedule_size(batch_size, num_heads_kv, aic_num, aiv_num))


def _parse_int_list(s):
    if s is None:
        return []
    s = str(s).strip()
    if not s:
        return []
    return [int(x.strip()) for x in s.split(",") if x.strip() != ""]


def _parse_shape(s):
    return tuple(_parse_int_list(s))


def _norm_dtype(v):
    if v is None:
        return None
    s = str(v).strip().upper()
    if s in ("FP4_E2M1", "FP4E2M1"):
        return "fp4_e2m1"
    if s in ("BF16",):
        return "bfloat16"
    if s in ("FLOAT8_E8M0",):
        return "uint8"
    return str(v).strip().lower()


def _compute_v_descale_shape(B, N_kv, V_D, skv, max_seqlen_kv=-1):
    s2 = max_seqlen_kv if max_seqlen_kv >= 0 else (max(skv) if skv else 0)
    n_blocks = (s2 + 31) // 32
    if n_blocks % 2 != 0:
        n_blocks += 1
    return (B, N_kv, n_blocks // 2, V_D, 2)


def _half_last_dim(shape):
    return tuple(list(shape[:-1]) + [shape[-1] // 2])


def _find_col_by_header(ws, header_name, start_col=1):
    for c in range(start_col, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v is not None and str(v).strip() == header_name:
            return c
    return None


def _build_col_map(ws):
    col_map = {}
    for key, hdr_name in HEADER.items():
        col = _find_col_by_header(ws, hdr_name)
        if col is not None:
            col_map[key] = col
    for hdr_name in OPTIONAL_HEADER:
        col = _find_col_by_header(ws, hdr_name)
        if col is not None:
            col_map[hdr_name] = col
    return col_map


def excel_row_to_csv_row(ws, row_idx, col_map, api_name, testcase_suffix):
    def g(key):
        col = col_map.get(key)
        return ws.cell(row=row_idx, column=col).value if col is not None else None

    def g_opt(key):
        col = col_map.get(key)
        return ws.cell(row=row_idx, column=col).value if col is not None else None

    def _raw_str(v):
        if v is None:
            return None
        s = str(v).strip()
        return s if s != "" else None

    name = str(g("name")).strip() + testcase_suffix
    B = int(g("batch_size"))
    N_q = int(g("num_heads_q"))
    N_kv = int(g("num_heads_kv"))
    D = int(g("head_dim"))
    G = N_q // N_kv if N_kv > 0 else 1
    V_D = D

    sq = _parse_int_list(g("seqused_q_value"))
    skv = _parse_int_list(g("seqused_kv_value"))
    cu_sq = _parse_int_list(g("cu_seqlens_q_value"))
    cu_skv = _parse_int_list(g("cu_seqlens_kv_value"))

    ms_q_raw = g("max_seqlen_q")
    ms_kv_raw = g("max_seqlen_kv")
    max_seqlen_q = (
        int(ms_q_raw) if ms_q_raw is not None and str(ms_q_raw).strip() != "" else -1
    )
    max_seqlen_kv = (
        int(ms_kv_raw) if ms_kv_raw is not None and str(ms_kv_raw).strip() != "" else -1
    )

    win_left = g("win_left")
    win_left_val = (
        2147483647 if (win_left is None or int(win_left) == -1) else int(win_left)
    )
    win_right = g("win_right")
    win_right_val = (
        2147483647 if (win_right is None or int(win_right) == -1) else int(win_right)
    )

    q_shape = _half_last_dim(_parse_shape(g("q_shape")))
    k_shape = _half_last_dim(_parse_shape(g("k_shape")))
    v_shape = _half_last_dim(_parse_shape(g("v_shape")))
    q_descale_shape = _parse_shape(g("q_descale_shape"))
    k_descale_shape = _parse_shape(g("k_descale_shape"))
    v_descale_shape = _parse_shape(g("v_descale_shape"))

    def _slot_shape(key):
        # 可选 tensor：Excel 空 → None（主算子收到 None 而非空 (0,) 张量，避免 segv）
        sh = _parse_shape(g(key))
        return sh if sh else None

    block_table_shape = _slot_shape("block_table_shape")
    p_scale_shape = _slot_shape("p_scale_shape")
    cu_seqlens_q_shape = _slot_shape("cu_seqlens_q_shape")
    cu_seqlens_kv_shape = _slot_shape("cu_seqlens_kv_shape")
    seqused_q_shape = (len(sq),) if sq else None
    seqused_kv_shape = (len(skv),) if skv else None
    sinks_shape = _slot_shape("learnable_sink_shape")
    attn_mask_shape = _slot_shape("attn_mask_shape")

    layout_q = str(g("layout_q")).strip()
    metadata_shape = _metadata_slot_shape(B, N_kv, sq, cu_sq, layout_q)

    shapes = [
        q_shape,
        k_shape,
        v_shape,
        q_descale_shape,
        k_descale_shape,
        v_descale_shape,
        block_table_shape,
        p_scale_shape,
        cu_seqlens_q_shape,
        cu_seqlens_kv_shape,
        seqused_q_shape,
        seqused_kv_shape,
        sinks_shape,
        attn_mask_shape,
        metadata_shape,
    ]
    tensor_view_shapes = repr(tuple(shapes))

    def _dt(key, default):
        return _norm_dtype(g(key)) or default

    dtypes = [
        "'uint8'",
        "'uint8'",
        "'uint8'",  # q/k/v packed
        "'uint8'",
        "'uint8'",
        "'uint8'",  # descale e8m0
        "'int32'",  # block_table
        "'float32'",  # p_scale
        "'int32'",
        "'int32'",
        "'int32'",
        "'int32'",  # cu_seqlens/seqused
        "'float32'",  # sinks
        "'int8'",  # attn_mask
        "'int32'",  # metadata
    ]
    tensor_dtypes = "(" + ",".join(dtypes) + ")"

    softmax_scale_raw = g("softmax_scale")
    if softmax_scale_raw is not None and str(softmax_scale_raw).strip() != "":
        softmax_scale = float(softmax_scale_raw)
    else:
        softmax_scale = 1.0 / math.sqrt(D)

    quant_mode = int(g("quant_mode")) if g("quant_mode") is not None else 5
    mask_mode = int(g("mask_mode")) if g("mask_mode") is not None else 0

    attrs = {
        "B": B,
        "N_q": N_q,
        "N_kv": N_kv,
        "G": G,
        "D": D,
        "V_D": V_D,
        "seqused_q_values": sq,
        "seqused_kv_values": skv,
        "cu_seqlens_q_values": cu_sq,
        "cu_seqlens_kv_values": cu_skv,
        "max_seqlen_q": max_seqlen_q,
        "max_seqlen_kv": max_seqlen_kv,
        "layout_q": layout_q,
        "layout_q_descale": str(g("layout_q_descale")).strip(),
        "layout_kv": str(g("layout_kv")).strip(),
        "layout_out": str(g("layout_out")).strip(),
        "kv_storage_mode": "continue",
        "block_size": 0,
        "q_dtype": _norm_dtype(g("q_dtype")) or "fp4_e2m1",
        "kv_dtype": _norm_dtype(g("k_dtype")) or "fp4_e2m1",
        "out_dtype": _norm_dtype(g("attn_out_dtype")) or "bfloat16",
        "quant_mode": quant_mode,
        "mask_mode": mask_mode,
        "win_left": win_left_val,
        "win_right": win_right_val,
        "return_softmax_lse": bool(g("return_softmax_lse"))
        if g("return_softmax_lse") is not None
        else False,
        "softmax_scale": softmax_scale,
        "inner_precise": 0,
        "device_id": 0,
        "graph_path": 0,
        "data_range_q": str(g("q_datarange")) if g("q_datarange") is not None else 1.0,
        "data_range_k": str(g("k_datarange")) if g("k_datarange") is not None else 1.0,
        "data_range_v": str(g("v_datarange")) if g("v_datarange") is not None else 1.0,
        "block_table_shape": list(block_table_shape),
        "block_table_dtype": _norm_dtype(g("block_table_dtype")),
        "p_scale_value": (
            float(g("p_scale_value"))
            if g("p_scale_value") is not None and str(g("p_scale_value")).strip() != ""
            else None
        ),
        "p_scale_shape": list(p_scale_shape),
        "p_scale_dtype": _norm_dtype(g("p_scale_dtype")),
        "p_scale_datarange": _raw_str(g("p_scale_datarange")),
        "sinks_shape": list(sinks_shape),
        "sinks_dtype": _norm_dtype(g("learnable_sink_dtype")),
        "attn_mask_shape": list(attn_mask_shape),
        "attn_mask_dtype": _norm_dtype(g("attn_mask_dtype")),
        "q_descale_dtype": _norm_dtype(g("q_descale_dtype")),
        "k_descale_dtype": _norm_dtype(g("k_descale_dtype")),
        "v_descale_dtype": _norm_dtype(g("v_descale_dtype")),
        "seqused_q_dtype": _norm_dtype(g("seqused_q_dtype")),
        "seqused_kv_dtype": _norm_dtype(g("seqused_kv_dtype")),
        "cu_seqlens_q_dtype": _norm_dtype(g("cu_seqlens_q_dtype")),
        "cu_seqlens_kv_dtype": _norm_dtype(g("cu_seqlens_kv_dtype")),
        "softmax_lse_dtype": _norm_dtype(g("softmax_lse_dtype")),
    }
    attributes = repr(attrs)

    pt_raw = g_opt("precision_tolerances")
    ap_raw = g_opt("absolute_precision")
    precision_tolerances = (
        str(pt_raw).strip() if pt_raw is not None else DEFAULT_PRECISION_TOLERANCES
    )
    absolute_precision = (
        str(ap_raw).strip() if ap_raw is not None else DEFAULT_ABSOLUTE_PRECISION
    )

    return [
        name,
        api_name,
        tensor_view_shapes,
        tensor_dtypes,
        "",
        attributes,
        "",
        "",
        precision_tolerances,
        absolute_precision,
    ]


def main():
    parser = argparse.ArgumentParser(description="Excel mxfp4 sheet -> TTK e2e CSV")
    parser.add_argument("--excel", default="B008QFA_红线用例.xlsx", help="Excel 路径")
    parser.add_argument("--sheet", default="mxfp4", help="sheet 名 (默认 mxfp4)")
    args = parser.parse_args()

    excel_path = (
        args.excel if os.path.isabs(args.excel) else os.path.join(_HERE, args.excel)
    )
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if args.sheet not in wb.sheetnames:
        raise ValueError(f"sheet '{args.sheet}' 不存在, 可用: {wb.sheetnames}")
    ws = wb[args.sheet]

    col_map = _build_col_map(ws)

    header = [
        "testcase_name",
        "api_name",
        "tensor_view_shapes",
        "tensor_dtypes",
        "tensor_formats",
        "attributes",
        "output_tensor_indexes",
        "golden_api",
        "precision_tolerances",
        "absolute_precision",
    ]
    name_col = col_map.get("name")
    data_row_indices = [
        r
        for r in range(2, ws.max_row + 1)
        if name_col and ws.cell(row=r, column=name_col).value is not None
    ]

    for fname, api_name, suffix in CSV_PROFILES:
        out_path = os.path.join(_HERE, fname)
        rows = [
            excel_row_to_csv_row(ws, r, col_map, api_name, suffix)
            for r in data_row_indices
        ]
        with open(out_path, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(header)
            writer.writerows(rows)
        print(f"-> wrote {out_path} ({len(rows)} cases, api_name={api_name})")


if __name__ == "__main__":
    main()
